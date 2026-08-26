"""
Enhancement script for 2026 Illinois CTP Business Plan
Applies 13 content recommendations to REVISED.docx → ENHANCED.docx.
All anchor searches are logged as APPLIED or ANCHOR NOT FOUND.
"""

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import copy, os

path_in  = os.path.expanduser(
    "~/champ-pm/BusinessPlan_2026/2026StateILJointBusinessPlan_REVISED.docx")
path_out = os.path.expanduser(
    "~/champ-pm/BusinessPlan_2026/2026StateILJointBusinessPlan_ENHANCED.docx")

doc = Document(path_in)
log = []

# ══════════════════════════════════════════════════════════════════════
# UTILITIES
# ══════════════════════════════════════════════════════════════════════

def full_text(para):
    return ''.join(r.text or '' for r in para.runs)

def replace_in_para(para, old, new):
    t = full_text(para)
    if old not in t:
        return False
    # Try per-run first
    for run in para.runs:
        if old in (run.text or ''):
            run.text = run.text.replace(old, new)
            return True
    # Multi-run fallback
    if para.runs:
        para.runs[0].text = t.replace(old, new)
        for r in para.runs[1:]:
            r.text = ''
    return True

def replace_all(doc, old, new):
    count = 0
    for para in doc.paragraphs:
        if replace_in_para(para, old, new):
            count += 1
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    if replace_in_para(para, old, new):
                        count += 1
    return count

def find_para(needle, occurrence=0):
    """Return the occurrence-th paragraph whose full text contains needle."""
    hits = [p for p in doc.paragraphs if needle in full_text(p)]
    if len(hits) > occurrence:
        return hits[occurrence]
    return None

def make_para_elem(text_content, style_val=None):
    """Build a w:p XML element with optional style."""
    new_p = OxmlElement('w:p')
    if style_val is not None:
        pPr = OxmlElement('w:pPr')
        pStyle = OxmlElement('w:pStyle')
        pStyle.set(qn('w:val'), style_val)
        pPr.append(pStyle)
        new_p.append(pPr)
    new_r = OxmlElement('w:r')
    new_t = OxmlElement('w:t')
    new_t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    new_t.text = text_content
    new_r.append(new_t)
    new_p.append(new_r)
    return new_p

def insert_after(ref_para, *para_elems):
    """
    Insert one or more w:p elements directly after ref_para, in order.
    Uses reversed addnext so first element ends up first.
    """
    ref = ref_para._element
    for elem in reversed(para_elems):
        ref.addnext(elem)

def para_index(para):
    paras = list(doc.paragraphs)
    try:
        return paras.index(para)
    except ValueError:
        return -1

# ══════════════════════════════════════════════════════════════════════
# REC 1 — Neutral text replacement
# ══════════════════════════════════════════════════════════════════════
n = replace_all(doc,
    "underserved populations",
    "communities with limited Risk MAP participation or technical resources")
n += replace_all(doc,
    "Potential strategies to enhance Community Outreach and Mitigation:",
    "Community Outreach and Mitigation \u2014 Program Strategies:")
if n:
    log.append(f"[REC 1]  APPLIED — neutral replacements: {n} change(s)")
else:
    log.append("[REC 1]  ANCHOR NOT FOUND — neutral replacement targets not found")

# ══════════════════════════════════════════════════════════════════════
# REC 2 — Community reach paragraph (neutral EJ framing)
# ══════════════════════════════════════════════════════════════════════
ANCHOR2 = ("Develop outreach materials that help community officials' outreach to their "
           "constituents to inform them of flood risk and ways to reduce their risk.")
ref = find_para(ANCHOR2)
if ref is None:
    # Try shorter needle
    ref = find_para("Develop outreach materials that help community officials")
if ref:
    p = make_para_elem(
        "A specific program priority for FFY2026 is extending flood risk information and CTP "
        "technical assistance to Illinois communities of all sizes and technical capacities. "
        "Many smaller municipalities, rural townships, and county governments with limited "
        "dedicated floodplain management staff have had limited engagement with the Risk MAP "
        "process, and may be unaware of available technical assistance resources, NFIP "
        "requirements, or flood risk products applicable to their area. ISWS and IDNR/OWR will "
        "develop targeted outreach and simplified guidance materials for these communities, "
        "coordinating outreach through county emergency management agencies, regional planning "
        "councils, and local government associations. Geographic analysis of communities with "
        "lower CRS participation, older FIRMs, or limited engagement in recent Risk MAP projects "
        "will help direct these efforts efficiently."
    )
    insert_after(ref, p)
    log.append("[REC 2]  APPLIED — community reach paragraph inserted")
else:
    log.append("[REC 2]  ANCHOR NOT FOUND — outreach materials paragraph not located")

# ══════════════════════════════════════════════════════════════════════
# REC 3 — Funding numbers placeholder (cost-benefit + numbers)
# ══════════════════════════════════════════════════════════════════════
ANCHOR3 = ("Program management funds are requested to maintain the illinoisfloodmaps.org "
           "website and for general program management.")
ref = find_para(ANCHOR3)
if ref is None:
    ref = find_para("illinoisfloodmaps.org website and for general program management")
if ref:
    p_numbers = make_para_elem(
        "[INSERT: Specify requested program management and COMS funding amounts for FFY2026. "
        "Cite current-year funding levels, requested increase amount or percentage, and "
        "supporting justification tied to expanded scope (e.g., 2D BLE initiative, increased "
        "COMS and MT-2 activity, community outreach growth).]"
    )
    p_cost = make_para_elem(
        "As a university-based Cooperating Technical Partner, ISWS provides substantial "
        "cost-effectiveness for the federal investment in Illinois flood hazard mapping. "
        "University overhead structures, established research infrastructure, and long-term "
        "institutional retention of technical expertise and project data allow ISWS to deliver "
        "hydrologic and hydraulic studies, flood mapping, community outreach, and COMS "
        "activities at rates significantly below those of equivalent private consultant work. "
        "The University of Illinois provides account management support, and all periodic "
        "program audits have been passed without negative findings \u2014 a record of financial "
        "accountability spanning more than 20 years of continuous CTP operation."
    )
    insert_after(ref, p_cost, p_numbers)
    log.append("[REC 3]  APPLIED — cost-benefit paragraph and funding placeholder inserted")
else:
    log.append("[REC 3]  ANCHOR NOT FOUND — program management funds paragraph not located")

# ══════════════════════════════════════════════════════════════════════
# REC 4/6/11 — Executive Summary outcomes and program year paragraphs
# ══════════════════════════════════════════════════════════════════════
ANCHOR4 = "help FEMA meet its national goals"
ref = find_para(ANCHOR4)
if ref:
    p_year = make_para_elem(
        "The FFY2026 program year (June 2026\u2013May 2027) marks the launch of a significant "
        "new initiative: the Statewide Two-Dimensional Base Level Engineering (2D BLE) project, "
        "which will \u2014 for the first time \u2014 develop modern 2D hydraulic modeling "
        "coverage for the entire State of Illinois. The Kishwaukee watershed mapping program "
        "simultaneously advances into Phase 7, with Physical Map Revisions anticipated for "
        "DeKalb, McHenry, and Kane Counties, completing the comprehensive multi-year watershed "
        "data development effort initiated in FY2020. Multiple county FIRM projects across "
        "southern and central Illinois advance through final map production and due process "
        "phases. IDNR/OWR continues an active regulatory, dam safety, and flood hazard "
        "mitigation program, with over $250\u202fmillion in active mitigation project requests "
        "statewide."
    )
    p_achieve = make_para_elem(
        "Over the past program period, the Illinois CTP partnership has delivered measurable "
        "results in flood hazard identification, mapping, and risk reduction. Illinois now has "
        "77 communities enrolled in FEMA\u2019s Community Rating System (CRS), providing an "
        "average 20% reduction in flood insurance premiums to policyholders statewide. Illinois "
        "has maintained a 100\u202f% community ordinance adoption rate following the issuance of "
        "updated FIRMs \u2014 a national benchmark that Illinois has upheld without exception "
        "since the program\u2019s inception, and that very few states consistently achieve. "
        "Statewide LiDAR data acquisition is now complete for all 102 Illinois counties, "
        "establishing the elevation data foundation for all current and future flood hazard "
        "mapping work. These results demonstrate that the Illinois CTP model consistently "
        "delivers strong, measurable returns on the federal investment in flood hazard "
        "mitigation and mapping."
    )
    insert_after(ref, p_achieve, p_year)
    log.append("[REC 4/6/11]  APPLIED — achievements paragraph + program year paragraph inserted")
else:
    log.append("[REC 4/6/11]  ANCHOR NOT FOUND — 'help FEMA meet its national goals' not located")

# ══════════════════════════════════════════════════════════════════════
# REC 5 — SAFR metrics placeholder
# ══════════════════════════════════════════════════════════════════════
ANCHOR5A = "visualization of structures at risk"
ANCHOR5B = "rich array of products"
ref = None
for p in doc.paragraphs:
    t = full_text(p)
    if ANCHOR5A in t and ANCHOR5B in t:
        ref = p
        break
if ref is None:
    # Try individual needle
    ref = find_para(ANCHOR5A)
if ref:
    p = make_para_elem(
        "[INSERT: SAFR website impact metrics \u2014 include the total number of structures in "
        "the database, annual unique site visitors, number of structure-specific assessments "
        "uploaded through Risk MAP projects, and any examples of communities that have used "
        "SAFR data for planning or mitigation decisions. These metrics directly demonstrate "
        "program impact to FEMA reviewers and should be updated annually in future business "
        "plans.]"
    )
    insert_after(ref, p)
    log.append("[REC 5]  APPLIED — SAFR metrics placeholder inserted")
else:
    log.append("[REC 5]  ANCHOR NOT FOUND — SAFR paragraph with both needles not located")

# ══════════════════════════════════════════════════════════════════════
# REC 7 — Post-2019 disaster response placeholder
# ══════════════════════════════════════════════════════════════════════
ANCHOR7 = "Illinois had a Federal Disaster Declaration following flooding during the fall or 2019"
ref = find_para(ANCHOR7)
if ref is None:
    ref = find_para("Federal Disaster Declaration following flooding during the fall")
if ref is None:
    ref = find_para("fall or 2019")
if ref:
    p = make_para_elem(
        "[INSERT: Post-2019 flood event response and community assistance activities "
        "(2020\u20132025). Has CHAMP or OWR conducted flood surveillance, high-water mark "
        "surveys, community contacts, damage assessments, or emergency coordination in response "
        "to significant flood events in Illinois since 2019? Documenting post-disaster "
        "activities demonstrates ongoing program responsiveness. Consider any federally declared "
        "disasters, major flooding on the Illinois, Mississippi, Wabash, or Kaskaskia Rivers, "
        "or other significant local flood events during this period.]"
    )
    insert_after(ref, p)
    log.append("[REC 7]  APPLIED — post-2019 disaster placeholder inserted")
else:
    log.append("[REC 7]  ANCHOR NOT FOUND — 2019 disaster declaration paragraph not located")

# ══════════════════════════════════════════════════════════════════════
# REC 8 — Sangamon River decision note
# ══════════════════════════════════════════════════════════════════════
ANCHOR8 = ("Although not proposed for funding this year, the main stem of the Upper Sangamon "
           "River has been the subject of several FEMA initiated studies")
ref = find_para(ANCHOR8)
if ref is None:
    ref = find_para("main stem of the Upper Sangamon River has been the subject of")
if ref is None:
    ref = find_para("Upper Sangamon River has been the subject of several FEMA initiated")
if ref:
    # Walk forward to find end of Sangamon section (Table 6 or "been established")
    paras = list(doc.paragraphs)
    idx = para_index(ref)
    last_sangamon = ref
    for i in range(idx + 1, min(idx + 6, len(paras))):
        t = full_text(paras[i]).strip()
        if "Table 6" in t or "been established" in t:
            last_sangamon = paras[i]
            break
        if t:  # non-empty paragraph that's part of the section
            last_sangamon = paras[i]
    p = make_para_elem(
        "[DECISION NEEDED: The Upper Sangamon River hydraulic assessment is based on a draft "
        "HEC-RAS model from 2013 \u2014 now more than 13 years old. This project description "
        "has appeared in multiple successive business plan cycles without advancing to funded "
        "work. Before the next submission, a decision is needed on whether to: (a) include "
        "Upper Sangamon in the FY27 or FY28 Five-Year Plan as a new Discovery and data "
        "development project with updated hydrology and scope; (b) formally defer to a "
        "specific future funding year with a documented rationale; or (c) remove from the "
        "plan until conditions and priorities warrant re-initiation. The current "
        "\u201cwill need review\u201d language should not carry forward to future submissions "
        "without an explicit resolution.]"
    )
    insert_after(last_sangamon, p)
    log.append("[REC 8]  APPLIED — Sangamon decision note inserted after last Sangamon paragraph")
else:
    log.append("[REC 8]  ANCHOR NOT FOUND — Upper Sangamon River paragraph not located")

# ══════════════════════════════════════════════════════════════════════
# REC 9 — Kishwaukee Phase 7 expansion
# ══════════════════════════════════════════════════════════════════════
ANCHOR9 = "PMR Dekalb"
ref = find_para(ANCHOR9)
if ref is None:
    ref = find_para("PMR DeKalb")
if ref:
    # Clear all existing runs and set new text
    for run in ref.runs:
        run.text = ''
    # Remove any existing w:r children entirely, then rebuild
    p_elem = ref._element
    # Remove all existing runs from XML
    for r in p_elem.findall(qn('w:r')):
        p_elem.remove(r)
    # Add new run with replacement text
    new_r = OxmlElement('w:r')
    new_t = OxmlElement('w:t')
    new_t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    new_t.text = (
        "Phase 7 (FY26) \u2014 With hydrologic and hydraulic data development complete for "
        "the full Kishwaukee HUC8 watershed, Phase 7 will advance the results of Phases 1\u20136 "
        "into Physical Map Revision (PMR) products for DeKalb, McHenry, and Kane Counties. "
        "Updated regulatory floodplain delineations will be developed for applicable stream "
        "reaches incorporating the watershed-scale hydraulic models and survey data collected "
        "in prior phases. Completion of Phase 7 represents the culmination of the comprehensive "
        "multi-year Kishwaukee watershed mapping program initiated in FY2020 and delivers "
        "updated, model-backed Flood Insurance Rate Maps for communities throughout the watershed."
    )
    new_r.append(new_t)
    p_elem.append(new_r)
    log.append("[REC 9]  APPLIED — Kishwaukee Phase 7 paragraph replaced")
else:
    log.append("[REC 9]  ANCHOR NOT FOUND — 'PMR Dekalb' paragraph not located")

# ══════════════════════════════════════════════════════════════════════
# REC 10 — Levee Mapping and Assessment subsection
#   Find "Five Year Plan List of Priorities" heading,
#   get the paragraph immediately BEFORE it, insert_after that para.
# ══════════════════════════════════════════════════════════════════════
ANCHOR10 = "Five Year Plan List of Priorities"
five_year_para = find_para(ANCHOR10)
if five_year_para:
    paras = list(doc.paragraphs)
    idx = para_index(five_year_para)
    if idx > 0:
        prev_para = paras[idx - 1]
        levee_body = make_para_elem(
            "Levee performance and residual flood risk are significant components of flood "
            "hazard management in Illinois. IDNR/OWR is pursuing a statewide levee assessment "
            "initiative using LiDAR-based survey data to determine overtopping elevations, "
            "characterize potential levee breach inundation extents, and assess the economic "
            "consequences of levee failure for communities protected by levees throughout the "
            "state. This work supports FEMA\u2019s National Levee Database (NLD), informs NFIP "
            "floodplain management decisions in levee-protected areas, and provides local "
            "communities with improved information for emergency planning, capital improvement "
            "prioritization, and mitigation decision-making. Levee overtopping and breach "
            "scenarios are developed in coordination with FEMA and USACE as applicable. "
            "[INSERT: Current status and timeline of the statewide levee mapping initiative "
            "\u2014 identify specific levee systems, river segments, or counties where "
            "assessment work is currently in progress or planned within the five-year plan "
            "period.]"
        )
        levee_heading = make_para_elem("Levee Mapping and Assessment", style_val="Heading2")
        insert_after(prev_para, levee_heading, levee_body)
        log.append("[REC 10]  APPLIED — Levee Mapping section inserted before Five Year Plan heading")
    else:
        log.append("[REC 10]  ANCHOR NOT FOUND — no paragraph before Five Year Plan heading")
else:
    log.append("[REC 10]  ANCHOR NOT FOUND — 'Five Year Plan List of Priorities' not located")

# ══════════════════════════════════════════════════════════════════════
# REC 11 — Section 4 introductory paragraph (add BEFORE anchor)
# ══════════════════════════════════════════════════════════════════════
ANCHOR11 = "Current outreach meetings and process"
ref = find_para(ANCHOR11)
if ref:
    intro = make_para_elem(
        "Community outreach and stakeholder engagement are central to the success of the "
        "Risk MAP program in Illinois. Effective outreach ensures that updated flood hazard "
        "information is understood and actively used by the communities it is intended to "
        "serve \u2014 from local elected officials and building departments to residents, "
        "realtors, lenders, and insurers. ISWS and IDNR/OWR approach community outreach as "
        "an integrated element of every project phase, from initial Discovery through FIRM "
        "adoption and beyond. The strategies described in this section reflect a commitment "
        "to making flood risk information accessible and actionable for all Illinois "
        "communities, and to building the local partnerships and institutional capacity "
        "needed for long-term flood risk reduction."
    )
    ref._element.addprevious(intro)
    log.append("[REC 11]  APPLIED — Section 4 intro paragraph inserted before outreach heading")
else:
    log.append("[REC 11]  ANCHOR NOT FOUND — 'Current outreach meetings and process' not located")

# ══════════════════════════════════════════════════════════════════════
# REC 12 — Precipitation frequency update note
# ══════════════════════════════════════════════════════════════════════
ANCHOR12A = "Maintain \u201cISWS Bulletin 75, Precipitation Frequency Study for Illinois\u201d"
ANCHOR12B = 'Maintain "ISWS Bulletin 75, Precipitation Frequency Study for Illinois"'
ref = find_para(ANCHOR12A)
if ref is None:
    ref = find_para(ANCHOR12B)
if ref is None:
    # Broader search
    for p in doc.paragraphs:
        t = full_text(p)
        if "Bulletin 75" in t and "Precipitation Frequency" in t:
            ref = p
            break
if ref:
    p = make_para_elem(
        "ISWS is currently evaluating methodological approaches to update the precipitation "
        "frequency analysis for Illinois, incorporating more recent long-term precipitation "
        "records and advancing the technical basis for flood frequency estimation. Updated "
        "precipitation frequency data directly affects the accuracy of hydrologic models, "
        "regulatory floodplain delineations, and engineering design standards throughout the "
        "state. Accurate, current precipitation frequency data is essential to ensuring that "
        "flood hazard maps reflect observed hydrologic conditions and support sound floodplain "
        "management decisions. [INSERT: Note any current research activities, partnership "
        "discussions with NOAA, USGS, or FEMA regarding precipitation frequency methodology "
        "updates, and whether an updated Bulletin 75 or successor study is currently planned "
        "or under development.]"
    )
    insert_after(ref, p)
    log.append("[REC 12]  APPLIED — precipitation frequency note inserted")
else:
    log.append("[REC 12]  ANCHOR NOT FOUND — Bulletin 75 / Precipitation Frequency paragraph not located")

# ══════════════════════════════════════════════════════════════════════
# REC 13 — Stakeholder participation section expansion
# ══════════════════════════════════════════════════════════════════════
ANCHOR13 = ("Serve on the following committees and groups to monitor and inform of "
            "Risk MAP program priorities")
ref = find_para(ANCHOR13)
if ref is None:
    ref = find_para("Serve on the following committees and groups to monitor")
if ref:
    p = make_para_elem(
        "Active participation in these organizations keeps the Illinois CTP program connected "
        "to national program developments, emerging technical standards, and best practices "
        "being implemented by peer CTPs across the country. These relationships also provide "
        "channels for the Illinois program to share its own tools and approaches \u2014 "
        "including SAFR, custom GIS scripts, and MT-2 review innovations \u2014 with the "
        "broader Risk MAP community."
    )
    insert_after(ref, p)
    log.append("[REC 13]  APPLIED — stakeholder section expansion inserted")
else:
    log.append("[REC 13]  ANCHOR NOT FOUND — 'Serve on the following committees' paragraph not located")

# ══════════════════════════════════════════════════════════════════════
# SAVE
# ══════════════════════════════════════════════════════════════════════
doc.save(path_out)

print("=" * 60)
print("ENHANCEMENT COMPLETE")
print("=" * 60)
for entry in log:
    print(entry)
print()
print(f"Saved: {path_out}")
